#!/usr/bin/env python3
"""
Palimpsest — a local, self-contained CAT tool built on the self-refine loop.

Runs entirely on your own machine. Your book text and your API keys never
leave your computer except for the calls to the translation API itself.

Setup:
    pip install -r requirements.txt
    Add your API key(s) from the app's Settings screen (or, for Anthropic
    only, put a key in key.txt in this folder for backward compatibility).

Run:
    Double-click start.command (Mac) or start.bat (Windows).
    Or: python server.py
"""

import os
import re
import json
import glob
import shutil
import datetime
import gzip
import pickle
import io
import zlib
import struct
import tempfile
from flask import Flask, request, jsonify, send_from_directory, send_file

import keys
import spending
import providers
import prefs
import segtypes

app = Flask(__name__, static_folder=".")

HERE = os.path.dirname(os.path.abspath(__file__))

# Books live in ~/.palimpsest/, next to keys.enc.json/prefs.json/spending.jsonl -
# not inside the app's own folder. An install that gets updated by replacing
# its own directory (git pull is fine, but a fresh ZIP extracted over a
# deleted old one is not) would otherwise take your books with it.
CONFIG_DIR = os.path.join(os.path.expanduser("~"), ".palimpsest")
PROJECTS_DIR = os.path.join(CONFIG_DIR, "projects")
os.makedirs(PROJECTS_DIR, exist_ok=True)

# One-time migration for installs that started before this moved: pull any
# books left in the old app-local projects/ folder over, but only if the new
# location is still empty - never overwrite anything already migrated.
_OLD_PROJECTS_DIR = os.path.join(HERE, "projects")
if os.path.isdir(_OLD_PROJECTS_DIR) and not os.listdir(PROJECTS_DIR):
    for _fname in os.listdir(_OLD_PROJECTS_DIR):
        if _fname.endswith(".json"):
            shutil.move(os.path.join(_OLD_PROJECTS_DIR, _fname), os.path.join(PROJECTS_DIR, _fname))



def clean_pending(max_age_hours=24):
    """Imports whose review screens were abandoned: projects/_pending/* (the PDF, the analysis, the questions and
    answers) older than a day. Returns the count of files removed."""
    d = os.path.join(PROJECTS_DIR, "_pending")
    if not os.path.isdir(d):
        return 0
    cutoff, n = datetime.datetime.now().timestamp() - max_age_hours * 3600, 0
    for name in os.listdir(d):
        path = os.path.join(d, name)
        if name.endswith((".pdf", ".state.pkl.gz", ".review.json")) and os.path.isfile(path) and os.path.getmtime(path) < cutoff:
            os.remove(path)
            n += 1
    return n


clean_pending()

DEFAULT_PROVIDER, DEFAULT_MODEL = "anthropic", "claude-opus-5-5"


def parse_model_key(s):
    """'provider::model_id' -> (provider, model_id), falling back to the default."""
    if s and "::" in s:
        provider, model = s.split("::", 1)
        if provider in providers.PROVIDERS and model in providers.PROVIDERS[provider]["models"]:
            return provider, model
    return DEFAULT_PROVIDER, DEFAULT_MODEL


# --------------------------------------------------------------------------
# File reading + segmentation
# --------------------------------------------------------------------------

def missing_library(feature, lib):
    """the error for a feature whose library is not installed: the server starts without the optional ones (PDF,
    OCR fallback, Word), and the launcher installs them on the next start with a connection"""
    return (f"{feature} için {lib} kütüphanesi eksik. Palimpsest'i internete bağlıyken başlatıcıyla (start.command / "
            "start.bat) yeniden aç; eksik kütüphaneyi kurar.")


def extract_text(filename, raw_bytes):
    ext = filename.lower().rsplit(".", 1)[-1]
    if ext in ("txt", "md"):
        return raw_bytes.decode("utf-8", errors="replace")
    if ext == "docx":
        import io
        try:
            from docx import Document
        except ImportError:
            raise ValueError(missing_library(".docx okuma", "python-docx"))
        doc = Document(io.BytesIO(raw_bytes))
        return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
    raise ValueError(f"Unsupported file type: .{ext}")


def segment(text):
    """Split on blank lines, then move each footnote right behind the first
    paragraph that cites it. Returns (segments, warnings)."""
    segs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    return segtypes.reorder_footnotes(segs)


# --------------------------------------------------------------------------
# One-time whole-book context pass (no caching needed — runs once)
# --------------------------------------------------------------------------

def condense_brief(provider, model, synopsis):
    """Compress the brief to minimum tokens without losing any instruction.
    Runs once at setup; saves tokens on every later translation call, since
    this brief rides along - cached - on every paragraph for the rest of
    the book. A fixed "half its length" target caps compression early; an
    absolute-minimum instruction keeps squeezing past that if the brief
    allows it, without touching what "no instruction lost" means."""
    if not synopsis or len(synopsis) < 200:
        return synopsis, None
    system = ("You compress translation briefs to the fewest words possible "
              "while preserving every instruction, nuance, and stylistic note. "
              "Telegraphic phrasing, sentence fragments, and dropped articles "
              "are fine. Compress the phrasing, never the content — if in doubt, "
              "keep the detail and cut a word elsewhere.")
    user = ("Rewrite this translation brief in the fewest words that lose zero "
            "instructions, nuances, or stylistic notes — no target length, just "
            "the true minimum. Keep every distinct instruction about voice, tone, "
            "register, POV, and style. Do not summarize away specifics. "
            "Output ONLY the condensed brief.\n\n" + synopsis)
    try:
        text, cost, tokens = providers.call_simple(provider, model, system, user, max_tokens=1024)
    except Exception:
        return synopsis, None, None
    # The call is paid for either way, so its cost is always returned. The rewrite replaces the
    # (already validated) brief only if it is complete and actually shorter.
    if tokens.get("truncated"):
        return synopsis, (cost, tokens), "Özet sıkıştırması yarıda kesildi; sıkıştırılmamış özet kullanıldı."
    if not text.strip() or len(text) >= len(synopsis):
        return synopsis, (cost, tokens), "Özet sıkıştırılamadı (sonuç kısalmadı); sıkıştırılmamış özet kullanıldı."
    return text, (cost, tokens), None


def build_context(provider, model, src, tgt, full_text):
    system = (
        "You are a senior literary translator preparing to translate a book. "
        "Before translating, you read the whole text and produce a working brief."
    )
    sample = full_text
    if len(full_text) > 120000:
        third = len(full_text) // 3
        sample = (full_text[:40000] + "\n\n[...]\n\n"
                  + full_text[third:third + 40000] + "\n\n[...]\n\n"
                  + full_text[-40000:])
    user = (
        f"This is a book to be translated from {src} into {tgt}. Read it and "
        f"produce two things, clearly separated:\n\n"
        f"=== SYNOPSIS ===\nA concise brief: genre, narrative voice, point of "
        f"view, tone, register, and stylistic features a translator must "
        f"preserve. 150-250 words.\n\n"
        f"=== TERM SHEET ===\n"
        f"A glossary of ONLY terms where a translation DECISION exists and "
        f"consistency across the book matters. Include: invented or culture-specific "
        f"terms, recurring motifs/refrains, titles and forms of address/honorifics, "
        f"terms with an established {tgt} convention, and names that change FORM in "
        f"{tgt} (transliteration, declension, vowel harmony). "
        f"EXCLUDE: any name or word that renders identically in {tgt} (do NOT list "
        f"X = X), one-off proper nouns that never recur, and anything mechanical. "
        f"If a term needs no decision, omit it entirely. "
        f"Reason specifically about the {src}->{tgt} direction and the book's genre. "
        f"Format one per line as `source = ` leaving the target BLANK for the "
        f"translator to decide, unless a {tgt} rendering is strongly conventional. "
        f"Each line holds exactly ONE source term (one spelling or variant) and at most ONE target: never join "
        f"variants with '/', ';' or 'or' - give every variant its own line. "
        f"Prefer a short, high-value list over an exhaustive one.\n\n"
        f"BOOK:\n{sample}"
    )
    return providers.call_simple(provider, model, system, user, max_tokens=4096)


SYNOPSIS_MAX_CHARS = 4000   # the prompt asks for 150-250 words (~1700 chars); this is the "absurd" line


def validate_context(blob, tokens):
    """Returns None if the model's context-pass answer is usable, else a short Turkish reason.
    A cut-off answer, a missing SYNOPSIS/TERM SHEET marker or an oversized synopsis is never
    allowed to reach the review screen."""
    if tokens.get("truncated"):
        return "model çıktısı uzunluk sınırında kesildi"
    if not re.search(r"===\s*SYNOPSIS\s*===", blob, re.I) or not re.search(r"===\s*TERM SHEET\s*===", blob, re.I):
        return "model beklenen biçimde yanıt vermedi"
    syn, _terms = split_context(blob)
    if not syn.strip():
        return "model boş bir özet döndürdü"
    if len(syn) > SYNOPSIS_MAX_CHARS:
        return "özet beklenenden çok uzun geldi"
    return None


def _fmt_tokens(n):
    return f"{int(n):,}".replace(",", ".")


def _fmt_cost(c):
    return ("$%.2f" % c if c >= 0.005 else "$%.4f" % c).replace(".", ",")


def _spent(tokens, cost):
    n = sum(tokens.get(k, 0) or 0 for k in ("input", "output", "cache_read", "cache_write"))
    return f"{_fmt_tokens(n)} token, {_fmt_cost(cost)} harcandı"


DEBUG_DIR = os.path.join(CONFIG_DIR, "debug")


def dump_debug(kind, raw, meta):
    """Local only, only on failure: keep the model's raw answer so the failure can be inspected."""
    os.makedirs(DEBUG_DIR, mode=0o700, exist_ok=True)
    path = os.path.join(DEBUG_DIR, datetime.datetime.now().strftime(f"{kind}-%Y%m%d-%H%M%S-%f.txt"))
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(f"{k}: {v}" for k, v in meta.items()) + "\n--- raw model output ---\n" + raw)
    os.chmod(path, 0o600)
    print(f"  ! setup validation failed ({meta.get('reason')}); raw model output saved to {path}", flush=True)
    return path


def split_context(blob):
    syn, terms = blob, ""
    m = re.search(r"===\s*TERM SHEET\s*===", blob, re.I)
    if m:
        syn, terms = blob[:m.start()], blob[m.end():]
    syn = re.sub(r"===\s*SYNOPSIS\s*===", "", syn, flags=re.I).strip()
    return syn, terms.strip()


# --------------------------------------------------------------------------
# Per-segment translation (single pass, with caching)
# --------------------------------------------------------------------------

def build_standing_context(synopsis, terms, rules, style_sample):
    """The part that stays IDENTICAL across calls — cacheable."""
    parts = []
    if synopsis: parts.append(f"PROJECT BRIEF:\n{synopsis}")
    if terms: parts.append(f"TERM SHEET (translate these consistently):\n{terms}")
    if rules: parts.append(f"RULES (follow strictly, no exceptions):\n{rules}")
    if style_sample: parts.append(f"STYLE TO MATCH:\n{style_sample}")
    return "\n\n".join(parts) if parts else "(no project context provided)"


def _edge_snippet(text, chars, from_end):
    """A short excerpt from a neighboring paragraph, for flow context only -
    it's never translated, so word-perfect boundaries don't matter, but this
    trims to the nearest whitespace instead of slicing mid-word. Sent uncached
    on every call, so full neighboring paragraphs (some run 1-2k chars) were
    pure waste for what's just meant to convey "what came right before/after"."""
    if len(text) <= chars:
        return text
    if from_end:
        cut = text[-chars:]
        sp = cut.find(" ")
        return ("…" + cut[sp + 1:]) if sp != -1 else ("…" + cut)
    cut = text[:chars]
    sp = cut.rfind(" ")
    return (cut[:sp] + "…") if sp != -1 else (cut + "…")


def _memory_text(memory):
    """memory arrives either as the old pre-joined string or as a list of
    [source, final] pairs. For pairs, both sides lose their markers and refs
    before being cut to 80 characters (cutting first could split a [^ref])."""
    if isinstance(memory, str):
        return memory
    lines = []
    for pair in memory or []:
        try:
            s, f = pair
        except (TypeError, ValueError):
            continue
        s, f = segtypes.strip_for_context(s or ""), segtypes.strip_for_context(f or "")
        if s and f:
            lines.append(s[:80] + " => " + f[:80])
    return "\n".join(lines)


def build_per_call(src, tgt, text, before, after, memory):
    """The part that changes per paragraph — NOT cached."""
    memory = _memory_text(memory)
    parts = []
    if memory: parts.append(f"EARLIER APPROVED TRANSLATIONS (match these):\n{memory}")
    if before: parts.append(f"PRECEDING TEXT (context only, do NOT translate):\n{_edge_snippet(before, 220, from_end=True)}")
    if after: parts.append(f"FOLLOWING TEXT (context only, do NOT translate):\n{_edge_snippet(after, 220, from_end=False)}")
    parts.append(
        f"Translate ONLY the SEGMENT below into {tgt}. "
        f"Follow the brief, term sheet, rules, and style. "
        f"Output ONLY the translation.\n\nSEGMENT:\n{text}"
    )
    return "\n\n".join(parts)


# --------------------------------------------------------------------------
# Routes
# --------------------------------------------------------------------------

@app.route("/")
def index():
    return send_from_directory(".", "index.html")


@app.route("/models")
def models():
    return jsonify(providers.list_catalog())


@app.route("/keys/status")
def keys_status():
    return jsonify(keys.key_status())


@app.route("/keys", methods=["POST"])
def keys_set():
    d = request.get_json()
    provider = d.get("provider")
    if provider not in providers.PROVIDERS:
        return jsonify({"error": "Unknown provider."}), 400
    keys.set_key(provider, d.get("api_key", ""))
    return jsonify({"ok": True, "has_key": bool(keys.get_key(provider))})


@app.route("/keys/clear", methods=["POST"])
def keys_clear():
    d = request.get_json()
    provider = d.get("provider")
    if provider not in providers.PROVIDERS:
        return jsonify({"error": "Unknown provider."}), 400
    keys.clear_key(provider)
    return jsonify({"ok": True})


@app.route("/spending")
def spending_summary():
    return jsonify(spending.summary())


@app.route("/prefs")
def prefs_get():
    return jsonify(prefs.get_all())


@app.route("/prefs", methods=["POST"])
def prefs_set():
    patch = request.get_json() or {}
    return jsonify(prefs.update(patch))


@app.route("/projects")
def list_projects():
    out = []
    for p in sorted(glob.glob(os.path.join(PROJECTS_DIR, "*.json"))):
        try:
            with open(p, encoding="utf-8") as f:
                d = json.load(f)
            done = sum(1 for x in d.get("finals", []) if x)
            out.append({"id": d.get("id"), "name": d.get("name", d.get("id")),
                        "count": len(d.get("segments", [])), "done": done,
                        "updated": d.get("updated", "")})
        except Exception:
            continue
    out.sort(key=lambda x: x.get("updated", ""), reverse=True)
    return jsonify(out)


@app.route("/project/<pid>")
def load_project(pid):
    p = project_path(pid)
    if not os.path.exists(p):
        return jsonify({"error": "Project not found."}), 404
    with open(p, encoding="utf-8") as f:
        return jsonify(json.load(f))


@app.route("/save", methods=["POST"])
def save_project():
    d = request.get_json()
    if not d.get("id"):
        return jsonify({"error": "missing id"}), 400
    d["updated"] = datetime.datetime.now().isoformat(timespec="seconds")
    with open(project_path(d["id"]), "w", encoding="utf-8") as f:
        json.dump(d, f, ensure_ascii=False, indent=2)
    if os.path.exists(pdf_path(d["id"], True)) and not os.path.exists(pdf_path(d["id"])):
        os.replace(pdf_path(d["id"], True), pdf_path(d["id"]))
    for ext in ("state.pkl.gz", "review.json"):          # the import is a project now: its review is done
        if os.path.exists(pending_file(d["id"], ext)):
            os.remove(pending_file(d["id"], ext))
    return jsonify({"ok": True, "updated": d["updated"]})


@app.route("/rename", methods=["POST"])
def rename_project():
    d = request.get_json()
    pid, name = d.get("id"), (d.get("name") or "").strip()
    if not pid or not name:
        return jsonify({"error": "missing id or name"}), 400
    p = project_path(pid)
    if not os.path.exists(p):
        return jsonify({"error": "Project not found."}), 404
    with open(p, encoding="utf-8") as f:
        proj = json.load(f)
    proj["name"] = name
    proj["updated"] = datetime.datetime.now().isoformat(timespec="seconds")
    with open(p, "w", encoding="utf-8") as f:
        json.dump(proj, f, ensure_ascii=False, indent=2)
    return jsonify({"ok": True})


@app.route("/duplicate", methods=["POST"])
def duplicate_project():
    d = request.get_json()
    pid = d.get("id")
    p = project_path(pid)
    if not os.path.exists(p):
        return jsonify({"error": "Project not found."}), 404
    with open(p, encoding="utf-8") as f:
        proj = json.load(f)
    proj["id"] = "p_" + str(int(datetime.datetime.now().timestamp() * 1000))
    proj["name"] = proj.get("name", "Untitled") + " (copy)"
    proj["updated"] = datetime.datetime.now().isoformat(timespec="seconds")
    with open(project_path(proj["id"]), "w", encoding="utf-8") as f:
        json.dump(proj, f, ensure_ascii=False, indent=2)
    if os.path.exists(pdf_path(pid)):                    # the copy keeps its page preview
        shutil.copyfile(pdf_path(pid), pdf_path(proj["id"]))
    return jsonify({"ok": True, "id": proj["id"]})


@app.route("/delete", methods=["POST"])
def delete_project():
    d = request.get_json()
    pid = d.get("id")
    p = project_path(pid)
    if not os.path.exists(p):
        return jsonify({"error": "Project not found."}), 404
    trash = os.path.join(PROJECTS_DIR, "_trash")
    os.makedirs(trash, exist_ok=True)
    os.replace(p, os.path.join(trash, os.path.basename(p)))
    if os.path.exists(pdf_path(pid)):                    # the PDF goes with its project
        os.replace(pdf_path(pid), os.path.join(trash, os.path.basename(pdf_path(pid))))
    return jsonify({"ok": True})


def _png(width, height, data, stride, channels):
    """raw RGB(A) rows -> PNG bytes (zlib only, no image library)"""
    rows = b"".join(b"\x00" + bytes(data[y * stride:y * stride + width * channels]) for y in range(height))
    chunk = lambda kind, body: struct.pack(">I", len(body)) + kind + body + struct.pack(">I", zlib.crc32(kind + body) & 0xFFFFFFFF)
    ihdr = struct.pack(">IIBBBBB", width, height, 8, 6 if channels == 4 else 2, 0, 0, 0)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", zlib.compress(rows, 6)) + chunk(b"IEND", b"")


@app.route("/page_image")
def page_image():
    """One page of a project's PDF copy, drawn on request at `scale` (1 = 72 dpi) - nothing is written to disk.
    The boxes in segMeta are PDF points from the top-left corner: multiply by the image width / page width."""
    pid, page = request.args.get("project", ""), request.args.get("page", "")
    path = (pdf_path(pid) if os.path.exists(pdf_path(pid)) else pdf_path(pid, True)) if pid else ""   # a review: still pending
    if not path or not os.path.exists(path):
        return jsonify({"error": "Bu projenin PDF kopyası yok."}), 404
    try:
        import pypdfium2 as pdfium
    except ImportError:
        return jsonify({"error": missing_library("Sayfa önizlemesi", "pypdfium2")}), 400
    try:
        scale = min(4.0, max(0.5, float(request.args.get("scale", "1.5"))))
        doc = pdfium.PdfDocument(path)
        n = int(page)
        if not 1 <= n <= len(doc):
            return jsonify({"error": f"Sayfa {n} yok (PDF {len(doc)} sayfa)."}), 404
        bm = doc[n - 1].render(scale=scale, rev_byteorder=True)
        channels = 4 if bm.mode in ("RGBA", "RGBX") else 3
        png = _png(bm.width, bm.height, bm.buffer, bm.stride, channels)
    except ValueError:
        return jsonify({"error": "Geçersiz sayfa ya da ölçek."}), 400
    return send_file(io.BytesIO(png), mimetype="image/png", max_age=3600)


PLACEHOLDER_LINE = re.compile(r"^!\[([^\]\r\n]*)\]\([^)\s]+\)$", re.M)


def docx_markdown(md):
    """The app's markdown as Pandoc should read it: a figure/table/index placeholder stays visible text
    ("[Görsel · s.17]") instead of an image Pandoc would try to load. A book with parts ("# Part One {.kisim}")
    gets a Word level for them: parts Heading 1, chapters Heading 2, every other heading one level down."""
    md = PLACEHOLDER_LINE.sub(lambda m: "\\[" + m.group(1) + "\\]", md)
    if not re.search(r"(?m)^#[ \t].*\{\.kisim\}[ \t]*$", md):
        return md
    return re.sub(r"(?m)^(#{1,5})([ \t].*)$", lambda m: m.group(0) if re.search(r"\{\.kisim\}[ \t]*$", m.group(2))
                  else "#" + m.group(1) + m.group(2), md)


def note_mark(label):
    """a note's mark as printed, from its label (index.html's fnNumber): "b2-13" -> 13, "b1-a" -> a, "b1-s2" -> †"""
    m = re.match(r"b\d+-([^-]+)", label)
    if m:
        s = re.fullmatch(r"s(\d+)", m.group(1))
        return "*†‡§‖¶"[(int(s.group(1)) - 1) % 6] * ((int(s.group(1)) - 1) // 6 + 1) if s else m.group(1)
    n = re.search(r"\d+$", label)
    return n.group(0) if n else label


def _runs(par, text):
    """inline **bold**, *italic* and footnote refs [^x] (as superscript) into python-docx runs"""
    for m in re.finditer(r"(\*\*\*[^*]+\*\*\*|\*\*[^*]+\*\*|\*[^*]+\*|\[\^[^\]\s]+\]|[^*\[]+|[*\[])", text):
        s = m.group(0)
        if s.startswith("***"):
            r = par.add_run(s[3:-3]); r.bold = r.italic = True
        elif s.startswith("**"):
            r = par.add_run(s[2:-2]); r.bold = True
        elif s.startswith("*") and len(s) > 1:
            r = par.add_run(s[1:-1]); r.italic = True
        elif s.startswith("[^"):
            r = par.add_run(note_mark(s[2:-1])); r.font.superscript = True
        else:
            par.add_run(s)


def _cell(cell, seg):
    """one segment into a table cell, styled like the Word export: headings, captions, placeholders, plain text"""
    import segtypes as st
    t = st.seg_type(seg)["type"]
    cell.paragraphs[0].text = ""
    par = cell.paragraphs[0]
    if t == "heading":
        m = st.HEADING.match(seg)
        par.style = "Heading %d" % min(3, len(m.group(1)))
        _runs(par, seg[m.end():])
    elif t == "caption":
        par.style = "Caption"
        _runs(par, st.strip_marker(seg))
    elif t in ("epigraph", "verse", "attribution", "box", "blockquote"):
        par.style = {"blockquote": "Block Text"}.get(t, st.FENCE_STYLE.get(t))
        for k, line in enumerate(st.strip_marker(seg).split("\n")):
            if k:
                par.add_run().add_break()
            _runs(par, line)
    elif t == "image":
        par.style = "Body Text"
        par.add_run("[" + seg[2:seg.index("]")] + "]")
    elif t == "table":
        par.style = "Body Text"
        par.add_run(seg)
    else:
        par.style = "Body Text"
        _runs(par, st.strip_marker(seg) if t == "footnote" else seg)


def redaction_docx(rows, ref):
    """the review copy: one two-column table, source | translation, one row per segment, on the reference styles. A
    row's "comment" (an in-book page reference) becomes a Word comment on its source cell."""
    import docx
    d = docx.Document(ref)
    body = d.element.body
    for el in list(body):
        if not el.tag.endswith("sectPr"):
            body.remove(el)
    table = d.add_table(rows=1, cols=2)
    table.style = "Table"
    for cell, text in zip(table.rows[0].cells, ("Kaynak", "Çeviri")):
        cell.paragraphs[0].text = text
        cell.paragraphs[0].runs[0].bold = True
    for r in rows:
        cells = table.add_row().cells
        _cell(cells[0], r.get("src", ""))
        _cell(cells[1], r.get("tgt", ""))
        runs = cells[0].paragraphs[0].runs
        if r.get("comment") and runs:
            d.add_comment(runs, text=r["comment"], author="Palimpsest", initials="P")
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


@app.route("/export_docx", methods=["POST"])
def export_docx():
    """Word export. mode "word": the translation through Pandoc with reference.docx (Heading 1-3, Body Text, First
    Paragraph, Caption, Epigraph, Block Text, Verse, Attribution, Box; tables as Word tables, footnotes as Word
    footnotes, Pandoc comment spans as Word comments). mode "redaksiyon": a two-column review
    table. The markdown and the rows come from the page (the same exportText the Markdown export uses)."""
    d = request.get_json() or {}
    ref = os.path.join(HERE, "reference.docx")
    name = re.sub(r"[^A-Za-z0-9_-]+", "_", d.get("name") or "ceviri")
    if d.get("mode") == "redaksiyon":
        try:
            import docx                                 # only to know it is installed
        except ImportError:
            return jsonify({"error": missing_library("Redaksiyon kopyası", "python-docx")}), 400
        data = redaction_docx(d.get("rows") or [], ref)
    else:
        try:
            import pypandoc
        except ImportError:
            return jsonify({"error": missing_library("Word dışa aktarma", "pypandoc_binary")}), 400
        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "out.docx")
            pypandoc.convert_text(docx_markdown(d.get("markdown") or ""), "docx", format="markdown",
                                  outputfile=out, extra_args=["--reference-doc=" + ref])
            data = open(out, "rb").read()
    return send_file(io.BytesIO(data), as_attachment=True, download_name=name + ".docx",
                     mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


# --------------------------------------------------------------------------
# PDF import, then the review screen (free), then /setup (paid)
# --------------------------------------------------------------------------

def pending_file(pid, ext):
    safe = re.sub(r"[^a-zA-Z0-9_-]", "_", pid)[:80]
    return os.path.join(PROJECTS_DIR, "_pending", f"{safe}.{ext}")


def _load_review(pid):
    try:
        return json.load(open(pending_file(pid, "review.json"), encoding="utf-8"))
    except (OSError, ValueError):
        return None


def _decisions(rv):
    """the review's answers (in question order) and its edits, as pdf_import decisions"""
    import review
    out = []
    for q in rv["questions"]:
        a = rv["answers"].get(q["id"])
        if a:
            out += review.decisions_for(q, a.get("key"), a.get("all", True))
    return out + list(rv.get("edits", []))


def _render_review(pid, rv):
    """the pending import rendered with the review's decisions: (markdown, warnings, seg_meta, seg_bids) as the
    project will get them, footnotes moved behind the paragraph citing them"""
    import pdf_import
    state = pickle.loads(gzip.decompress(open(pending_file(pid, "state.pkl.gz"), "rb").read()))
    r = pdf_import.render(state, _decisions(rv))
    raw = [x.strip() for x in re.split(r"\n\s*\n", r["md"]) if x.strip()]
    segs, _ = segment(r["md"])
    return dict(md=r["md"], warnings=r["warnings"], raw_meta=r["meta"], meta=carry_meta(raw, segs, r["meta"]),
                bids=carry_meta(raw, segs, r["doc"].seg_bids), segs=segs, state=state)


def _review_payload(pid, rv, x=None):
    x = x or _render_review(pid, rv)
    segs, meta, bids, warnings, d = x["segs"], x["meta"], x["bids"], x["warnings"], x["state"]["doc"]
    import review
    return dict(rv, segments=segs, segMeta=meta, segBids=bids, warnings=warnings, integrity=review.integrity(x["state"]),
                pageSizes={p["p"]: [round(p["W"], 1), round(p["H"], 1)] for p in d.pages})


@app.route("/import", methods=["POST"])
def import_pdf():
    """A PDF's free half: analysis, the markdown as the importer reads it, and the review screen's questions. The
    PDF, the analysis and the review wait in projects/_pending/ until the project is saved (or a day passes)."""
    f, pid = request.files.get("file"), request.form.get("project_id", "")
    if not f or not pid or not f.filename.lower().endswith(".pdf"):
        return jsonify({"error": "Kontrol ekranı için bir PDF ve proje kimliği gerekli."}), 400
    try:
        import pdf_import
        import review
    except ImportError:
        return jsonify({"error": missing_library("PDF içe aktarma", "pypdfium2")}), 400
    data = f.read()
    try:
        state = pdf_import.analyse(data, request.form.get("first_page"), request.form.get("last_page"))
    except pdf_import.PdfImportError as e:
        return jsonify({"error": str(e)}), 400
    os.makedirs(os.path.dirname(pending_file(pid, "x")), exist_ok=True)
    with open(pending_file(pid, "state.pkl.gz"), "wb") as out:
        out.write(gzip.compress(pickle.dumps(state, protocol=pickle.HIGHEST_PROTOCOL), 3))
    rendered = pdf_import.render(state)              # the pristine state is on disk; this copy may change
    rv = dict(id=pid, name=request.form.get("name", ""), first=state["first"], last=state["last"],
              pages=len(state["doc"].pages), questions=review.questions(state, rendered), answers={}, edits=[],
              integrity=review.integrity(state))
    json.dump(rv, open(pending_file(pid, "review.json"), "w", encoding="utf-8"), ensure_ascii=False)
    keep_pending_pdf(pid, data)
    return jsonify(_review_payload(pid, rv))


@app.route("/import/<pid>", methods=["GET"])
def import_review(pid):
    """the review screen as it was left (a page reload)"""
    rv = _load_review(pid)
    if rv is None:
        return jsonify({"error": "Bu içe aktarma bulunamadı (bir günden eski olabilir); PDF'i yeniden yükle."}), 404
    return jsonify(_review_payload(pid, rv))


@app.route("/import/<pid>/lines", methods=["GET"])
def import_lines(pid):
    """the lines of a segment's blocks, for "böl": [{bid, line, text}] (line 0 of a later block splits between blocks)"""
    rv = _load_review(pid)
    if rv is None:
        return jsonify({"error": "Bu içe aktarma bulunamadı."}), 404
    x = _render_review(pid, rv)
    try:
        bids = x["bids"][int(request.args.get("seg", ""))] or []
    except (ValueError, IndexError):
        return jsonify({"error": "Böyle bir paragraf yok."}), 400
    blocks = {b.get("bid"): b for p in x["state"]["doc"].pages for b in p["blocks"]}
    out = [dict(bid=bid, line=k, text=l["plain"].replace("\x02", "-").strip())
           for bid in bids if bid in blocks for k, l in enumerate(blocks[bid]["lines"])]
    return jsonify({"lines": out})


@app.route("/import/<pid>/answers", methods=["POST"])
def import_answers(pid):
    """answers {question id: {key, all}} and edits (pdf_import decisions: join/split/drop) saved, the import rendered
    again with them"""
    rv = _load_review(pid)
    if rv is None:
        return jsonify({"error": "Bu içe aktarma bulunamadı (bir günden eski olabilir); PDF'i yeniden yükle."}), 404
    d = request.get_json() or {}
    rv["answers"] = {k: v for k, v in (d.get("answers") or {}).items() if v and v.get("key") not in (None, "skip")}
    rv["edits"] = [e for e in (d.get("edits") or []) if isinstance(e, dict)]
    x = _render_review(pid, rv)
    # a block the user made a note that found no callout is asked "which callout?" right after the question that
    # made it one; such a question stays while answered (its answer is a decision), goes when the block is no note
    import review
    notes = {b.get("bid") for p in x["state"]["doc"].pages for b in p["blocks"] if b.get("user_note")}
    new = [q for q in review.questions(x["state"], dict(meta=x["raw_meta"], md=x["md"])) if q["type"] == "bagla"]
    have = {q["id"] for q in rv["questions"]}
    qs = [q for q in rv["questions"] if q["type"] != "bagla" or q["bid"] in notes and (
        q["id"] in rv["answers"] or any(n["id"] == q["id"] for n in new))]
    for n in (n for n in new if n["id"] not in have):
        at = next((k + 1 for k, q in enumerate(qs) if n["bid"] in q.get("bids", []) or any(
            n["bid"] in (t if isinstance(t, list) else [t]) for t in q["targets"])), len(qs))
        qs.insert(at, n)
    rv["questions"] = qs                             # no new answer among the changes: the render stands
    json.dump(rv, open(pending_file(pid, "review.json"), "w", encoding="utf-8"), ensure_ascii=False)
    return jsonify(_review_payload(pid, rv, x))


@app.route("/setup", methods=["POST"])
def setup():
    f = request.files.get("file")
    src = request.form.get("source_lang", "English")
    tgt = request.form.get("target_lang", "Turkish")
    project_name = request.form.get("name", "")
    provider, model = parse_model_key(request.form.get("model"))
    import_id = request.form.get("import_id", "")
    if not f and not import_id:
        return jsonify({"error": "No file received."}), 400
    pdf_warnings, pdf_scanned, pdf_meta, pdf_bytes = [], False, None, None
    try:
        if import_id:                                   # a reviewed PDF: rendered with the review's decisions
            rv = _load_review(import_id)
            if rv is None:
                return jsonify({"error": "Bu içe aktarma bulunamadı (bir günden eski olabilir); PDF'i yeniden yükle."}), 404
            import pdf_import
            x = _render_review(import_id, rv)
            text, pdf_warnings, pdf_meta = x["md"], x["warnings"], x["raw_meta"]
            pdf_scanned = pdf_import.SCANNED_WARNING in pdf_warnings
        elif f.filename.lower().endswith(".pdf"):
            # imported here, not at the top: a missing pypdfium2 must only switch PDF import off,
            # never stop the server from starting
            try:
                import pdf_import
            except ImportError:
                return jsonify({"error": missing_library("PDF içe aktarma", "pypdfium2")}), 400
            try:
                pdf_bytes = f.read()
                text, pdf_warnings, pdf_meta = pdf_import.convert_with_meta(
                    pdf_bytes, request.form.get("first_page"), request.form.get("last_page"))
                pdf_scanned = pdf_import.SCANNED_WARNING in pdf_warnings
            except pdf_import.PdfImportError as e:
                return jsonify({"error": str(e)}), 400
        else:
            text = extract_text(f.filename, f.read())
    except ValueError as e:
        return jsonify({"error": str(e)}), 400          # unsupported type, a missing library: said as it is
    except Exception as e:
        return jsonify({"error": f"Could not read the file: {e}"}), 400
    raw_segs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    segs, import_warnings = segment(text)
    seg_meta = carry_meta(raw_segs, segs, pdf_meta) if pdf_meta is not None else None
    if not segs:
        return jsonify({"error": "No text found in the file."}), 400
    # The model reads the book without markdown markers or footnote refs; the
    # stored segments keep them.
    text = "\n\n".join(segtypes.strip_for_context(s) for s in segs)
    if not keys.get_key(provider):
        return jsonify({"error": f"No API key for {provider}. Add one from Settings."}), 500
    spent_tokens, spent_cost = {}, 0.0
    try:
        blob, cost1, tokens1 = build_context(provider, model, src, tgt, text)
        # paid for already: log it before anything can reject the answer
        spending.log_call("", project_name, provider, model, "off", tokens1, cost1)
        spent_tokens, spent_cost = tokens1, cost1
        problem = validate_context(blob, tokens1)
        if problem:
            dump_debug("setup", blob, {"reason": problem, "provider": provider, "model": model,
                                       "stop_reason": tokens1.get("stop_reason"), "output_tokens": tokens1.get("output"),
                                       "cost": cost1, "book_chars": len(text)})
            return jsonify({"error": f"Özet üretilemedi: {problem} ({_spent(tokens1, cost1)}). Tekrar deneyebilirsin."}), 502
        synopsis, terms = split_context(blob)
        terms = split_term_variants(terms)
        # Condense the brief once now, to cut tokens on every later call.
        condensed, condense_info, condense_note = condense_brief(provider, model, synopsis)
        total_cost = cost1
        if condense_info:
            cost2, tokens2 = condense_info
            spending.log_call("", project_name, provider, model, "off", tokens2, cost2)
            total_cost += cost2
        warnings = pdf_warnings + [w["message"] for w in import_warnings] + ([condense_note] if condense_note else [])
        return jsonify({"segments": segs, "synopsis": condensed, "terms": terms,
                        "setup_cost": round(total_cost, 4), "count": len(segs),
                        "warnings": warnings, "scanned": pdf_scanned, "segMeta": seg_meta,
                        "pdfKept": keep_pending_pdf(request.form.get("project_id"), pdf_bytes)
                        or bool(import_id and os.path.exists(pdf_path(import_id, True)))})
    except Exception as e:
        spent = f" ({_spent(spent_tokens, spent_cost)})" if spent_tokens else ""
        return jsonify({"error": f"Bağlam okuması başarısız: {e}{spent}"}), 500


@app.route("/clean_glossary", methods=["POST"])
def clean_glossary():
    """Re-filter an existing glossary: drop X=X identities, one-offs, and noise.
    Driven by language pair and genre. Returns a cleaned term list."""
    d = request.get_json()
    provider, model = parse_model_key(d.get("model"))
    src, tgt = d.get("source_lang", "English"), d.get("target_lang", "Turkish")
    terms = (d.get("terms") or "").strip()
    synopsis = (d.get("synopsis") or "").strip()
    project_id, project_name = d.get("id", ""), d.get("name", "")
    if not terms:
        return jsonify({"terms": "", "cost": 0})
    if not keys.get_key(provider):
        return jsonify({"error": f"No API key for {provider}. Add one from Settings."}), 500
    try:
        system = (
            "You are a translation glossary editor. You ruthlessly prune glossaries "
            "down to only entries that represent a real, consistency-critical "
            f"translation decision in the {src}->{tgt} direction."
        )
        user = (
            f"Book context (for genre/register):\n{synopsis or '(none given)'}\n\n"
            f"Here is a raw {src}->{tgt} glossary. Clean it:\n"
            f"- REMOVE any entry that renders identically in {tgt} (X = X).\n"
            f"- REMOVE one-off proper nouns with no recurring significance.\n"
            f"- REMOVE mechanical or trivial entries.\n"
            f"- KEEP invented terms, motifs, honorifics/forms of address, "
            f"culture-specific terms, and names that change FORM in {tgt}.\n"
            f"- Preserve any target renderings already filled in.\n"
            f"- Sort alphabetically by source term.\n"
            f"Output ONLY the cleaned glossary, one entry per line as `source = target` "
            f"(target may be blank). No commentary.\n\n"
            f"GLOSSARY:\n{terms}"
        )
        text, cost, tokens = providers.call_simple(provider, model, system, user, max_tokens=2048)
        spending.log_call(project_id, project_name, provider, model, "off", tokens, cost)
        return jsonify({"terms": text.strip(), "cost": round(cost, 4)})
    except Exception as e:
        return jsonify({"error": f"Cleanup failed: {e}"}), 500


@app.route("/translate_simple", methods=["POST"])
def translate_simple():
    """Single-pass translation with prompt caching."""
    d = request.get_json()
    provider, model = parse_model_key(d.get("model"))
    src, tgt = d.get("source_lang", "English"), d.get("target_lang", "Turkish")
    raw_text = d.get("text", "").strip()
    src_type = segtypes.seg_type(raw_text)
    text, ref_labels = segtypes.to_prompt(raw_text)
    text = text.strip()
    effort = d.get("effort", "off")
    project_id, project_name = d.get("id", ""), d.get("name", "")
    if not text:
        return jsonify({"error": "Empty segment."}), 400
    if not keys.get_key(provider):
        return jsonify({"error": f"No API key for {provider}. Add one from Settings."}), 500
    try:
        system_text = (
            f"You are a professional literary translator from {src} into {tgt}. "
            f"Translate within the book's context. Follow all rules and the term sheet exactly. "
            f"Markers like \u27e61\u27e7 stand for footnote references: keep each one exactly as "
            f"written, right after the word it belongs to."
        )
        if src_type["type"] == "table":
            # a markdown table: the structure must come back unchanged (from_model compares rows and columns)
            system_text += (" The text to translate is a markdown table: translate only the text inside the cells. "
                            "Keep every | character, every row, the |---| separator row and the empty cells exactly "
                            "as they are; do not add or remove rows or columns.")
        if src_type["type"] == "verse" or (src_type["type"] == "epigraph" and "\n" in text):
            # a poem: from_model warns when its lines come back as one
            system_text += (" The text to translate is verse: keep it in lines, one line of the translation per "
                            "line break; the number of lines may differ where the translation needs it, but never "
                            "run the poem together into a single line.")
        standing = build_standing_context(
            d.get("synopsis", ""), d.get("terms", ""),
            d.get("rules", ""), segtypes.strip_block(d.get("style_sample", "") or ""))
        per_call = build_per_call(
            src, tgt, text,
            segtypes.strip_for_context(d.get("before", "") or ""),
            segtypes.strip_for_context(d.get("after", "") or ""),
            d.get("memory", ""))

        result, cost, tokens = providers.call(
            provider, model, system_text, standing, per_call, effort=effort)
        spending.log_call(project_id, project_name, provider, model, effort, tokens, cost)

        final, warnings = segtypes.from_model(result, text, ref_labels, src_type)
        return jsonify({
            "final": final, "cost": cost,
            "tokens": tokens, "warnings": warnings
        })
    except Exception as e:
        return jsonify({"error": f"Translation failed: {e}"}), 500


VARIANT_SEP = re.compile(r"\s*/\s*|\s*;\s*|\s+or\s+", re.I)


def split_term_variants(terms):
    """One variant per glossary line, whatever the model sent: "colour / color = renk" becomes two lines, both "= renk";
    "A / B = X / Y" pairs up when both sides have the same number of variants. Other lines pass through untouched."""
    out = []
    for line in (terms or "").split("\n"):
        if "=" not in line:
            out.append(line)
            continue
        src, tgt = line.split("=", 1)
        srcs = [s.strip() for s in VARIANT_SEP.split(src.strip()) if s.strip()]
        tgts = [s.strip() for s in VARIANT_SEP.split(tgt.strip()) if s.strip()]
        if len(srcs) <= 1:
            out.append(line)
            continue
        for k, s in enumerate(srcs):
            t = tgts[k] if len(tgts) == len(srcs) else tgt.strip()
            out.append(f"{s} = {t}".rstrip())
    return "\n".join(out)


def carry_meta(raw_segs, segs, meta):
    """segMeta follows its segment through segment()'s footnote reordering (matched by text, in order); a count
    mismatch means the two splits disagree, and then there is no meta at all rather than misplaced meta."""
    if len(meta) != len(raw_segs):
        return None
    queue = {}
    for s, m in zip(raw_segs, meta):
        queue.setdefault(s, []).append(m)
    return [queue[s].pop(0) if queue.get(s) else None for s in segs]


def pdf_path(pid, pending=False):
    safe = re.sub(r"[^a-zA-Z0-9_-]", "_", pid)[:80]
    return os.path.join(PROJECTS_DIR, "_pending" if pending else "", safe + ".pdf")


def keep_pending_pdf(pid, data):
    """The imported PDF waits under projects/_pending/ until the project is first saved (the review screen can still
    be abandoned); /save moves it next to the project as <id>.pdf, for the page preview."""
    if not pid or not data:
        return False
    os.makedirs(os.path.dirname(pdf_path(pid, True)), exist_ok=True)
    with open(pdf_path(pid, True), "wb") as f:
        f.write(data)
    return True


def project_path(pid):
    safe = re.sub(r"[^a-zA-Z0-9_-]", "_", pid)[:80]
    return os.path.join(PROJECTS_DIR, safe + ".json")


if __name__ == "__main__":
    import webbrowser, threading
    if not any(keys.key_status().values()):
        print("\n  ! No API key found for any provider.")
        print("  ! Open the app and add one from the Settings screen,")
        print("  ! or (Anthropic only) put a key in key.txt and restart.\n")
    else:
        print("\n  Palimpsest is running.")
        print("  Open  http://localhost:5001  in your browser (opening it for you now).\n")
        threading.Timer(1.2, lambda: webbrowser.open("http://localhost:5001")).start()
    app.run(port=5001, debug=False)
