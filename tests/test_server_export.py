"""Run: python3 -I tests/test_server_export.py   (plain asserts; a temporary HOME, no network, no real projects)

Page images for the source preview, the Word exports (Pandoc with reference.docx; the two-column review copy) and the
glossary variant split of new setups."""
import io
import os
import sys
import struct
import zipfile
import tempfile

HOME = tempfile.mkdtemp(prefix="palimpsest-test-")
os.environ["HOME"] = HOME
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))
import server
import docx
import pypdfium2 as pdfium

failed = 0


def check(name, cond, extra=""):
    global failed
    print(("PASS " if cond else "FAIL ") + name + ("" if cond else "  " + str(extra)))
    if not cond:
        failed += 1


c = server.app.test_client()

# ---------------------------------------------------------------- page images
pdf = pdfium.PdfDocument.new()
pdf.new_page(300, 400)
pdf.new_page(300, 400)
buf = io.BytesIO()
pdf.save(buf)
os.makedirs(server.PROJECTS_DIR, exist_ok=True)
open(server.pdf_path("p_img"), "wb").write(buf.getvalue())
before = sorted(os.listdir(server.PROJECTS_DIR))
r = c.get("/page_image?project=p_img&page=2&scale=2")
w, h = struct.unpack(">II", r.data[16:24])
check("page image: a PNG of the page at the asked scale", r.status_code == 200 and r.mimetype == "image/png" and (w, h) == (600, 800), (r.status_code, w, h))
r = c.get("/page_image?project=p_img&page=1&scale=99")
check("page image: scale is capped", struct.unpack(">II", r.data[16:24]) == (1200, 1600))
check("page image: nothing is written to disk", sorted(os.listdir(server.PROJECTS_DIR)) == before)
check("page image: a page that does not exist is 404, garbage is 400, no PDF is 404",
      c.get("/page_image?project=p_img&page=3").status_code == 404 and c.get("/page_image?project=p_img&page=x").status_code == 400
      and c.get("/page_image?project=nope&page=1").status_code == 404)

# ---------------------------------------------------------------- Word: the translation through Pandoc
md = "\n\n".join([
    "# Birinci Bölüm",
    "İlk paragraf burada[^n1] ve *italik* bir sözcük.",
    "[^n1]: Bu bir dipnottur.",
    "İkinci paragraf **kalın** ile.",
    "![Görsel · s.17](gorsel-17-1)",
    '::: {custom-style="Caption"}\n*Şekil 1:* Bir kedi\n:::',
    "| Ad | Not |\n|---|---|\n| a \\| b | bir |",
    "## Alt başlık",
    "[UNTRANSLATED] Henüz çevrilmedi.",
])
r = c.post("/export_docx", json={"mode": "word", "name": "deneme", "markdown": md})
check("Word export answers a .docx", r.status_code == 200 and r.data[:2] == b"PK" and "deneme.docx" in r.headers.get("Content-Disposition", ""))
d = docx.Document(io.BytesIO(r.data))
by_text = {p.text: p.style.name for p in d.paragraphs if p.text.strip()}
check("styles: # -> Heading 1, ## -> Heading 2", by_text.get("Birinci Bölüm") == "Heading 1" and by_text.get("Alt başlık") == "Heading 2", by_text)
first = next(t for t in by_text if t.startswith("İlk paragraf"))
check("styles: first paragraph after a heading is First Paragraph, the next Body Text",
      by_text[first] == "First Paragraph" and by_text.get("İkinci paragraf kalın ile.") == "Body Text", by_text)
check("styles: the caption fence becomes the Caption style", by_text.get("Şekil 1: Bir kedi") == "Caption", by_text)
check("a figure placeholder stays visible text", "[Görsel · s.17]" in by_text, list(by_text))
check("untranslated segments keep their [UNTRANSLATED] mark", any(t.startswith("[UNTRANSLATED]") for t in by_text))
check("a markdown table becomes a Word table (escaped pipe inside the cell)", len(d.tables) == 1 and d.tables[0].cell(1, 0).text == "a | b",
      [[c.text for c in r.cells] for r in d.tables[0].rows] if d.tables else None)
notes = zipfile.ZipFile(io.BytesIO(r.data)).read("word/footnotes.xml").decode("utf-8")
check("footnotes become Word footnotes", "Bu bir dipnottur." in notes)
styles = {s.name for s in d.styles}
check("reference.docx carries the contract styles", {"Heading 1", "Heading 2", "Heading 3", "Body Text", "First Paragraph", "Caption"} <= styles)

# ---------------------------------------------------------------- Word: the two-column review copy
rows = [{"src": "# Chapter", "tgt": "# Bölüm"}, {"src": "A line with *one* note[^1].", "tgt": "Bir *not*[^1] satırı."},
        {"src": "![Görsel · s.4](gorsel-4-1)", "tgt": "![Görsel · s.4](gorsel-4-1)"},
        {"src": "Not yet.", "tgt": "[UNTRANSLATED] Not yet."}]
r = c.post("/export_docx", json={"mode": "redaksiyon", "name": "r", "rows": rows})
d = docx.Document(io.BytesIO(r.data))
t = d.tables[0] if d.tables else None
check("review copy: one two-column table, a header row and one row per segment", t is not None and len(t.columns) == 2 and len(t.rows) == 5)
check("review copy: header Kaynak | Çeviri", [x.text for x in t.rows[0].cells] == ["Kaynak", "Çeviri"])
check("review copy: headings keep their style inside the cells", t.cell(1, 1).paragraphs[0].style.name == "Heading 1" and t.cell(1, 1).text == "Bölüm")
check("review copy: inline italics and footnote refs as runs", any(r_.italic and r_.text == "not" for r_ in t.cell(2, 1).paragraphs[0].runs)
      and any(r_.font.superscript and r_.text == "1" for r_ in t.cell(2, 1).paragraphs[0].runs))
check("review copy: placeholders as text, untranslated marked", t.cell(3, 0).text == "[Görsel · s.4]" and t.cell(4, 1).text.startswith("[UNTRANSLATED]"))

# ---------------------------------------------------------------- no Pandoc: a clear Turkish error, nothing else breaks
saved = sys.modules.get("pypandoc")
sys.modules["pypandoc"] = None
r = c.post("/export_docx", json={"mode": "word", "markdown": "x"})
check("without pypandoc: 400 with the install hint", r.status_code == 400 and "pypandoc_binary" in r.get_json()["error"])
r = c.post("/export_docx", json={"mode": "redaksiyon", "rows": rows})
check("the review copy does not need Pandoc", r.status_code == 200)
if saved is not None:
    sys.modules["pypandoc"] = saved
else:
    del sys.modules["pypandoc"]

# ---------------------------------------------------------------- glossary: one variant per line
out = server.split_term_variants("colour / color = renk\nA / B = X / Y\nkoan; kōan =\ntrees or forests = orman\nsingle = tek\nno equals line")
check("joined variants become one line each (shared or paired targets)", out.split("\n") == [
    "colour = renk", "color = renk", "A = X", "B = Y", "koan =", "kōan =", "trees = orman", "forests = orman", "single = tek", "no equals line"], out)
import inspect
check("the setup prompt asks for one variant per line", "never join " in inspect.getsource(server.build_context)
      and "own line" in inspect.getsource(server.build_context))

import shutil
shutil.rmtree(HOME, ignore_errors=True)
print("FAILED %d" % failed if failed else "ALL PASS")
sys.exit(1 if failed else 0)
